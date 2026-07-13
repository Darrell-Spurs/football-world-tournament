#!/usr/bin/env python3
"""Local server for the World Cup Draft squad dashboard.

Serves the project folder (so viewer/squad_viewer_live.html can fetch the
files in resources/ live) AND exposes a small write API so the dashboard can
save player swaps and finalize squads back into resources/squad_registry.xlsx.

Endpoints
  GET  /...                     static files (no-cache)
  POST /api/save                {team, squad:[{id,role}], reserves:[{id,rid}]}
                                apply the new squad/reserve arrangement
  POST /api/finalize            {team}
                                delete the 9 reserves from the team sheet

Every write first copies squad_registry.xlsx into resources/backups/.

Run:  python serve.py     (or: py serve.py)
Stop: Ctrl+C
"""
import http.server
import socketserver
import json
import os
import sys
import shutil
import datetime
import webbrowser
import unicodedata

PORT = int(os.environ.get("PORT", "8000"))
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # project root
REGISTRY = os.path.join(ROOT, "resources", "squad_registry.xlsx")
BACKUP_DIR = os.path.join(ROOT, "resources", "backups")
os.chdir(ROOT)

try:
    import openpyxl
    EDIT_OK = True
except Exception:
    openpyxl = None
    EDIT_OK = False

# Column layout (1-based) in every team sheet and in All_Players
SQUAD_HEADER = ["Country", "Player", "PlayerID", "Position", "PositionDetail", "RAV",
                "MarketValue", "Club", "Age", "Nationality", "Second Nationality",
                "Nationality Source", "SquadRole"]
RES_HEADER = ["ReserveID", "Player", "PlayerID", "Position", "PositionDetail", "RAV",
              "MarketValue", "Club", "Age", "Nationality", "Second Nationality"]


def norm(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return "".join(ch for ch in s.lower() if ch.isalnum())


def backup():
    os.makedirs(BACKUP_DIR, exist_ok=True)
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    dst = os.path.join(BACKUP_DIR, f"squad_registry_{ts}.xlsx")
    shutil.copy2(REGISTRY, dst)
    return dst


def find_team_sheet(wb, team):
    for name in wb.sheetnames:
        if norm(name) == norm(team):
            return wb[name]
    return None


def read_players(ws):
    """Return (squad_rows, reserve_rows, reserve_header_row). Each row is a dict
    of column-name -> value plus _row (the sheet row index)."""
    squad, reserves = [], []
    mode = "squad"
    reserve_header_row = None
    for i in range(2, ws.max_row + 1):
        a = ws.cell(row=i, column=1).value
        if mode == "squad":
            if a is None or str(a).strip() == "":
                mode = "gap"
                continue
            rec = {SQUAD_HEADER[c]: ws.cell(row=i, column=c + 1).value for c in range(len(SQUAD_HEADER))}
            rec["_row"] = i
            squad.append(rec)
        elif mode == "gap":
            if a == "ReserveID":
                reserve_header_row = i
                mode = "reserve"
        elif mode == "reserve":
            if a is None or str(a).strip() == "":
                continue
            rec = {RES_HEADER[c]: ws.cell(row=i, column=c + 1).value for c in range(len(RES_HEADER))}
            rec["_row"] = i
            reserves.append(rec)
    return squad, reserves, reserve_header_row


def derive_source(country, nationality, second):
    if norm(nationality) == norm(country):
        return "First"
    if second and norm(second) == norm(country):
        return "Second"
    return "Others"


def apply_save(team, want_squad, want_reserves):
    """want_squad: [{id, role}] (26), want_reserves: [{id, rid}] (up to 9)."""
    wb = openpyxl.load_workbook(REGISTRY)
    ws = find_team_sheet(wb, team)
    if ws is None:
        raise ValueError(f"Team sheet not found: {team}")
    squad, reserves, res_hdr = read_players(ws)
    if res_hdr is None:
        raise ValueError("No reserve section on this sheet (already finalized?).")

    country = squad[0]["Country"] if squad else team
    # Full-attribute lookup by PlayerID across current squad + reserves
    attr = {}
    for r in squad:
        attr[str(r["PlayerID"])] = dict(r)
    for r in reserves:
        attr[str(r["PlayerID"])] = dict(r)

    def full_attrs(pid):
        a = attr.get(str(pid))
        if a is None:
            raise ValueError(f"Unknown PlayerID in this team: {pid}")
        return a

    squad_start = squad[0]["_row"]        # first squad data row
    res_start = res_hdr + 1               # first reserve data row

    # ---- rewrite squad rows ----
    for offset, item in enumerate(want_squad):
        a = full_attrs(item["id"])
        row = squad_start + offset
        vals = {
            "Country": country, "Player": a.get("Player"), "PlayerID": a.get("PlayerID"),
            "Position": a.get("Position"), "PositionDetail": a.get("PositionDetail"),
            "RAV": a.get("RAV"), "MarketValue": a.get("MarketValue"), "Club": a.get("Club"),
            "Age": a.get("Age"), "Nationality": a.get("Nationality"),
            "Second Nationality": a.get("Second Nationality"),
            "Nationality Source": a.get("Nationality Source") or
                derive_source(country, a.get("Nationality"), a.get("Second Nationality")),
            "SquadRole": item["role"],
        }
        for c, key in enumerate(SQUAD_HEADER):
            ws.cell(row=row, column=c + 1, value=vals[key])

    # ---- rewrite reserve rows ----
    for offset, item in enumerate(want_reserves):
        a = full_attrs(item["id"])
        row = res_start + offset
        vals = {
            "ReserveID": item["rid"], "Player": a.get("Player"), "PlayerID": a.get("PlayerID"),
            "Position": a.get("Position"), "PositionDetail": a.get("PositionDetail"),
            "RAV": a.get("RAV"), "MarketValue": a.get("MarketValue"), "Club": a.get("Club"),
            "Age": a.get("Age"), "Nationality": a.get("Nationality"),
            "Second Nationality": a.get("Second Nationality"),
        }
        for c, key in enumerate(RES_HEADER):
            ws.cell(row=row, column=c + 1, value=vals[key])

    # ---- keep All_Players in sync (it holds the 26-man squad only) ----
    if "All_Players" in wb.sheetnames:
        ap = wb["All_Players"]
        role_by_id = {str(it["id"]): it["role"] for it in want_squad}
        del_rows = [i for i in range(2, ap.max_row + 1)
                    if norm(ap.cell(row=i, column=1).value) == norm(country)]
        for i in reversed(del_rows):
            ap.delete_rows(i, 1)
        for it in want_squad:
            a = full_attrs(it["id"])
            ap.append([
                country, a.get("Player"), a.get("PlayerID"), a.get("Position"),
                a.get("PositionDetail"), a.get("RAV"), a.get("MarketValue"), a.get("Club"),
                a.get("Age"), a.get("Nationality"), a.get("Second Nationality"),
                a.get("Nationality Source") or
                    derive_source(country, a.get("Nationality"), a.get("Second Nationality")),
                it["role"],
            ])

    bpath = backup()
    wb.save(REGISTRY)
    return {"ok": True, "backup": os.path.basename(bpath)}


def apply_finalize(team):
    wb = openpyxl.load_workbook(REGISTRY)
    ws = find_team_sheet(wb, team)
    if ws is None:
        raise ValueError(f"Team sheet not found: {team}")
    squad, reserves, res_hdr = read_players(ws)
    if res_hdr is None:
        return {"ok": True, "note": "already finalized"}
    last_squad_row = squad[-1]["_row"] if squad else 1
    first_del = last_squad_row + 1
    last_del = ws.max_row
    if last_del >= first_del:
        ws.delete_rows(first_del, last_del - first_del + 1)
    bpath = backup()
    wb.save(REGISTRY)
    return {"ok": True, "removed_reserves": len(reserves), "backup": os.path.basename(bpath)}


class Handler(http.server.SimpleHTTPRequestHandler):
    def end_headers(self):
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        self.send_header("Expires", "0")
        super().end_headers()

    def log_message(self, fmt, *args):
        pass

    def _json(self, code, obj):
        body = json.dumps(obj).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):
        if not self.path.startswith("/api/"):
            self._json(404, {"ok": False, "error": "unknown endpoint"})
            return
        if not EDIT_OK:
            self._json(503, {"ok": False, "error": "openpyxl not installed on the server. Run: pip install openpyxl"})
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length) or b"{}")
        except Exception as e:
            self._json(400, {"ok": False, "error": f"bad request body: {e}"})
            return
        try:
            if self.path == "/api/save":
                res = apply_save(payload["team"], payload["squad"], payload.get("reserves", []))
            elif self.path == "/api/finalize":
                res = apply_finalize(payload["team"])
            else:
                self._json(404, {"ok": False, "error": "unknown endpoint"})
                return
            self._json(200, res)
        except PermissionError:
            self._json(423, {"ok": False, "error": "squad_registry.xlsx is locked. Close it in Excel and try again."})
        except Exception as e:
            self._json(500, {"ok": False, "error": str(e)})


def main():
    url = f"http://localhost:{PORT}/viewer/squad_viewer_live.html"
    print("=" * 60)
    print(" World Cup Draft — live squad dashboard (read/write)")
    print("=" * 60)
    print(f" Serving : {ROOT}")
    print(f" Open    : {url}")
    print(f" Editing : {'ENABLED' if EDIT_OK else 'DISABLED - run: pip install openpyxl'}")
    print(" Stop    : Ctrl+C")
    print("=" * 60)
    try:
        webbrowser.open(url)
    except Exception:
        pass
    try:
        with socketserver.TCPServer(("127.0.0.1", PORT), Handler) as httpd:
            httpd.serve_forever()
    except KeyboardInterrupt:
        print("\nStopped.")
    except OSError as e:
        print(f"\nCould not start server on port {PORT}: {e}")
        print("Set a different port, e.g.  set PORT=8010 && python serve.py")
        sys.exit(1)


if __name__ == "__main__":
    main()
